"""Step 01: resolve and validate the experiment ground truth.

For each experiment: expand config/experiments.yaml into a per-well role table,
normalise + patch the legacy platemap, normalise the timepoint file, and check
that config and platemap agree (lines, conditions, uniqueness, S.A. wells per
line). Also documents the XDP10 duplicate-block correction and the info-workbook
XDP8/XDP10 label swap as evidence tables.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from lib import cfg  # noqa: E402
from lib.paths import LOCAL_OUT, RAW_CACHE  # noqa: E402
from lib.provenance import StepRecorder, sha256_file  # noqa: E402

STEP = "01_metadata"
EXPECTED_SA_PER_LINE = {"XDP1": 49, "XDP3": 12, "XDP5": 6, "XDP8": 0, "XDP10": 2, "XDP11": 12}
# Documented exceptions where the legacy platemap Drug column is known to be incomplete.
DRUG_EXCEPTIONS = {"XDP1": "whole plate dosed with S.A.; platemap Drug column is blank for every well"}


def drug_to_condition(drug: str) -> str:
    if drug == "SA":
        return "sa"
    return "untreated" if drug == "" else "sirna"


def xlsx_blocks(path: Path) -> pd.DataFrame:
    """Summarise each experiment block of the info workbook: label, rows used, lines, S.A. note."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    blocks, cur = [], None
    for r in range(1, ws.max_row + 1):
        title = ws.cell(r, 1).value
        if title:
            cur = {"label": str(title).strip(), "note": ws.cell(r, 28).value, "lines": set(),
                   "rows_with_cells": set(), "sa_cols": set()}
            blocks.append(cur)
        if cur is None:
            continue
        row_letter = ws.cell(r, 2).value
        for c in range(3, 27):
            v = ws.cell(r, c).value
            if v is None:
                continue
            v = str(v).strip()
            if v == "S.A.":
                cur["sa_cols"].add(c - 2)
            elif "siRNA" not in v and row_letter:
                cur["lines"].add(v)
                cur["rows_with_cells"].add(str(row_letter))
    out = []
    for b in blocks:
        rows = sorted(b["rows_with_cells"])
        out.append({"workbook_label": b["label"], "lines": ", ".join(sorted(b["lines"])),
                    "n_lines": len(b["lines"]), "rows": f"{rows[0]}-{rows[-1]}" if rows else "",
                    "sa_columns": ", ".join(map(str, sorted(b["sa_cols"]))), "note": b["note"]})
    return pd.DataFrame(out)


def main() -> None:
    out = LOCAL_OUT / STEP
    rec = StepRecorder(STEP, out)
    scope = cfg.lines_in_scope()
    summary = {}
    all_roles = []
    for exp, e in cfg.experiments().items():
        rec.add_input(RAW_CACHE / e["platemap"], sha256_file(RAW_CACHE / e["platemap"]))
        rec.add_input(RAW_CACHE / e["timepoints"], sha256_file(RAW_CACHE / e["timepoints"]))
        roles = cfg.well_roles(exp)
        roles.insert(0, "exp", exp)
        all_roles.append(roles)
        pm_raw = cfg.read_platemap(exp, apply_patches=False)
        pm = cfg.read_platemap(exp)
        tps = cfg.read_timepoints(exp)
        tps.to_csv(out / f"timepoints_{exp}.csv", index=False)
        pm.to_csv(out / f"platemap_corrected_{exp}.csv", index=False)

        rec.check(f"{exp}.platemap_unique", f"{exp}: corrected platemap has one row per well",
                  pm.well.is_unique, {"rows": len(pm), "unique": pm.well.nunique()})
        m = roles.merge(pm, on="well", how="outer", indicator=True)
        rec.check(f"{exp}.wells_match", f"{exp}: config wells == platemap wells",
                  (m._merge == "both").all(),
                  {"only_config": m.loc[m._merge == "left_only", "well"].tolist(),
                   "only_platemap": m.loc[m._merge == "right_only", "well"].tolist()})
        both = m[m._merge == "both"].copy()
        rec.check(f"{exp}.lines_match", f"{exp}: config line == platemap line for every well",
                  (both.line == both.Sci_SampleID).all(),
                  both.loc[both.line != both.Sci_SampleID, ["well", "line", "Sci_SampleID"]].to_dict("records"))
        both["pm_condition"] = both.Drug.map(drug_to_condition)
        mism = both[both.pm_condition != both.condition]
        if exp in DRUG_EXCEPTIONS:
            rec.check(f"{exp}.conditions_match", f"{exp}: config condition vs platemap Drug "
                      f"(documented exception: {DRUG_EXCEPTIONS[exp]})", None, {"mismatched_wells": len(mism)})
        else:
            rec.check(f"{exp}.conditions_match", f"{exp}: config condition == platemap Drug for every well",
                      len(mism) == 0, mism[["well", "condition", "Drug"]].to_dict("records")[:20])
        sa_per_line = roles[(roles.condition == "sa") & roles.in_scope].groupby("line").size()
        exp_n = EXPECTED_SA_PER_LINE[exp]
        rec.check(f"{exp}.sa_wells_per_line", f"{exp}: S.A. wells per in-scope line == {exp_n}",
                  bool((sa_per_line == exp_n).all()) if exp_n else len(sa_per_line) == 0,
                  sa_per_line.to_dict())
        rec.check(f"{exp}.timepoints_monotone", f"{exp}: timepoint hours strictly increase",
                  tps.hours.is_monotonic_increasing and tps.hours.is_unique, tps.hours.tolist())

        per_line = (roles[roles.in_scope].groupby(["line", "condition"]).size()
                    .unstack(fill_value=0).reset_index())
        summary[exp] = {
            "n_wells": len(roles), "lines_in_scope": sorted(roles[roles.in_scope].line.unique()),
            "wells_by_line_condition": per_line.to_dict("records"),
            "timepoints": tps.tp.tolist(), "hours": tps.hours.tolist(),
            "platemap_rows_raw": len(pm_raw), "platemap_unique_wells_raw": int(pm_raw.well.nunique()),
            "sa": e["sa"],
        }

        if exp == "XDP10":
            # Evidence table: the appended block vs the first block for the same wells.
            first = pm_raw.iloc[:210].set_index("well")
            app = pm_raw.iloc[210:].copy()
            app["first_block_line"] = app.well.map(first["Sci_SampleID"])
            app["first_block_drug"] = app.well.map(first["Drug"])
            app["in_first_block"] = app.well.isin(first.index)
            app = app.rename(columns={"Sci_SampleID": "appended_line", "Drug": "appended_drug"})
            cols = ["well", "first_block_line", "first_block_drug", "appended_line", "appended_drug",
                    "in_first_block"]
            app[cols].to_csv(out / "xdp10_platemap_duplicate_block.csv", index=False)
            summary[exp]["duplicate_block"] = {
                "rows": len(app), "wells_also_in_first_block": int(app.in_first_block.sum()),
                "label_differs": int((app.in_first_block & (app.appended_line != app.first_block_line)).sum()),
                "sa_labels_in_appended_block": int((app.appended_drug == "SA").sum()),
            }

    pd.concat(all_roles).to_csv(out / "well_roles_all.csv", index=False)

    # Info workbook: which block is which experiment?
    xlsx = RAW_CACHE / "XDP-experiments-info.xlsx"
    rec.add_input(xlsx, sha256_file(xlsx))
    blocks = xlsx_blocks(xlsx)
    exp_lines = {exp: set(e["lines"]) for exp, e in cfg.experiments().items()}
    blocks["matches_data_of"] = blocks.lines.map(
        lambda s: ", ".join(exp for exp, ls in exp_lines.items()
                            if {x.strip() for x in s.split(",")} == {l.split("-", 1)[1] for l in ls}))
    blocks.to_csv(out / "info_workbook_blocks.csv", index=False)
    summary["info_workbook_blocks"] = blocks.to_dict("records")
    rec.summary = summary
    rec.write()
    print(f"{STEP}: {len(rec.checks)} checks, {len(rec.failed)} FAIL")
    for c in rec.checks:
        if c["status"] != "PASS":
            print("  ", c["status"], c["id"], c["value"])


if __name__ == "__main__":
    main()
